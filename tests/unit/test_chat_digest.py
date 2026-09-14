"""Reading a created file back: the facts its overview is written from."""

from __future__ import annotations

import io

from app.chat_artifacts.digest import describe


def xlsx(sheets: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestSpreadsheet:
    def test_rows_columns_and_ranges_are_counted(self):
        text = describe(xlsx({"Presidents": [
            ["Name", "Start", "End"],
            ["Washington", 1789, 1797],
            ["Adams", 1797, 1801],
            ["Jefferson", 1801, 1809],
        ]}), "excel")
        assert "3 data rows, 3 columns" in text
        assert "Start: numbers from 1789 to 1801" in text
        assert "Washington" in text

    def test_a_summary_sheet_reads_as_notes(self):
        text = describe(xlsx({
            "Summary": [["US Presidents", None], ["source", "general knowledge"]],
            "Data": [["A", "B", "C"], ["1", "2", "3"]],
        }), "excel")
        assert "2 sheet(s): Summary, Data" in text
        assert "Sheet 'Summary' (notes)" in text and "- source: general knowledge" in text

    def test_a_titled_sheet_with_any_name_reads_as_notes(self):
        text = describe(xlsx({"Sheet1": [["Report title", None], ["date", "2026-09-14"]]}), "excel")
        assert "(notes)" in text

    def test_a_small_two_column_table_stays_data(self):
        rows = [["Title", "Studio"], ["Belle", "Studio Chizu"], ["Suzume", "CoMix Wave"]]
        text = describe(xlsx({"Movies": rows}), "excel")
        assert "2 data rows, 2 columns" in text and "(notes)" not in text

    def test_few_distinct_values_are_listed(self):
        rows = [["Title", "Type"]] + [[f"t{i}", "Movie" if i % 2 else "Series"] for i in range(10)]
        assert "Type: 2 distinct values (Movie, Series)" in describe(xlsx({"S": rows}), "excel")

    def test_many_distinct_values_are_only_counted(self):
        rows = [["Title"]] + [[f"title {i}"] for i in range(30)] + [["x"]]
        text = describe(xlsx({"S": [r + ["", ""] for r in rows]}), "excel")
        assert "Title: 31 distinct values" in text


class TestOtherFormats:
    def test_csv(self):
        text = describe("Title,Year\nBelle,2021\nSuzume,2022\n".encode("utf-8-sig"), "csv")
        assert "2 data rows, 2 columns" in text and "Year: numbers from 2021 to 2022" in text

    def test_word_headings_are_marked(self):
        from docx import Document

        doc = Document()
        doc.add_heading("Phase 1", level=1)
        doc.add_paragraph("Plan the work.")
        buf = io.BytesIO()
        doc.save(buf)
        text = describe(buf.getvalue(), "word")
        assert "# Phase 1" in text and "Plan the work." in text

    def test_unreadable_bytes_never_raise(self):
        assert "could not be read back" in describe(b"not a workbook", "excel")
