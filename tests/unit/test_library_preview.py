"""The in-app viewer's previews: spreadsheets as grids, Word as paragraphs."""

from __future__ import annotations

import io
from datetime import datetime

import pytest

from app.library.preview import PreviewError, build_preview


def xlsx(sheets: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestSpreadsheet:
    def test_every_sheet_becomes_a_grid(self):
        p = build_preview("book.xlsx", xlsx({
            "Dishes": [["Name", "Price", "Halal"], ["Laksa", 8.5, True], ["Satay", 12.0, False]],
            "Notes": [["Prepared", datetime(2026, 9, 14)]],
        }))
        assert p["type"] == "table" and [s["name"] for s in p["sheets"]] == ["Dishes", "Notes"]
        dishes = p["sheets"][0]
        assert dishes["rows"][1] == ["Laksa", "8.5", "TRUE"]
        assert dishes["rows"][2] == ["Satay", "12", "FALSE"]
        assert p["sheets"][1]["rows"][0] == ["Prepared", "2026-09-14"]

    def test_every_row_of_a_big_sheet_is_sent(self):
        rows = [["n", "double"]] + [[i, i * 2] for i in range(3200)]
        sheet = build_preview("big.xlsx", xlsx({"S": rows}))["sheets"][0]
        assert len(sheet["rows"]) == 3201 and sheet["total_rows"] == 3201
        assert sheet["rows"][-1] == ["3199", "6398"] and sheet["truncated"] is False

    def test_every_column_is_sent(self):
        wide = [[f"c{i}" for i in range(77)], list(range(77))]
        sheet = build_preview("wide.xlsx", xlsx({"S": wide}))["sheets"][0]
        assert sheet["total_cols"] == 77 and len(sheet["rows"][0]) == 77

    def test_a_sheet_past_the_cell_budget_is_cut_and_says_so(self, monkeypatch):
        import app.library.preview as preview

        monkeypatch.setattr(preview, "MAX_CELLS", 100)
        rows = [["a", "b"]] + [[i, i] for i in range(200)]
        sheet = build_preview("big.xlsx", xlsx({"S": rows}))["sheets"][0]
        assert len(sheet["rows"]) == 50 and sheet["total_rows"] == 201 and sheet["truncated"] is True

    def test_trailing_empty_rows_are_not_counted(self):
        sheet = build_preview("t.xlsx", xlsx({"S": [["a", "b"], ["1", "2"], [None, None], [None]]}))["sheets"][0]
        assert sheet["total_rows"] == 2 and sheet["total_cols"] == 2 and sheet["truncated"] is False


class TestCsv:
    def test_commas(self):
        sheet = build_preview("m.csv", "Title,Year\nBelle,2021\n".encode("utf-8-sig"))["sheets"][0]
        assert sheet["name"] == "m" and sheet["rows"] == [["Title", "Year"], ["Belle", "2021"]]

    def test_semicolons_are_detected(self):
        sheet = build_preview("eu.csv", b"Name;Price\nKopi;1,80\nTeh;1,60\n")["sheets"][0]
        assert sheet["rows"][1] == ["Kopi", "1,80"]


class TestWord:
    def test_headings_paragraphs_and_tables(self):
        from docx import Document

        d = Document()
        d.add_heading("Phase 1", level=1)
        d.add_paragraph("Plan the work.")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Step", "Owner"
        t.cell(1, 0).text, t.cell(1, 1).text = "Design", "Tan"
        buf = io.BytesIO()
        d.save(buf)
        p = build_preview("plan.docx", buf.getvalue())
        assert p["type"] == "document"
        assert p["blocks"][0] == {"kind": "heading", "text": "Phase 1"}
        assert p["blocks"][1] == {"kind": "paragraph", "text": "Plan the work."}
        assert p["tables"][0]["rows"] == [["Step", "Owner"], ["Design", "Tan"]]


class TestOther:
    def test_a_pdf_is_left_to_the_browser(self):
        assert build_preview("report.pdf", b"%PDF-1.7")["type"] == "unsupported"

    def test_a_damaged_workbook_is_an_error_not_a_crash(self):
        with pytest.raises(PreviewError):
            build_preview("broken.xlsx", b"not a workbook")

    def test_too_large(self):
        assert build_preview("huge.xlsx", b"0" * (26 * 1024 * 1024))["type"] == "too_large"
