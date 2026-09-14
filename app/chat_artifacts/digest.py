"""What is actually inside a file — read back from its bytes.

The note that accompanies a created file is written from this, never from
what the builder was asked for: a model describing its own intentions would
describe a file that may not exist. Counts, ranges and names are computed
here, so the note can quote them without the model counting anything.
"""

from __future__ import annotations

import csv
import io
import re

_TEXT_CHARS = 4000
_SAMPLE_ROWS = 6
_MAX_SHEETS = 6
_MAX_DISTINCT = 12
_NUMBER = re.compile(r"-?\d+(?:\.\d+)?")


def describe(data: bytes, fmt: str) -> str:
    """A plain-text account of the file's contents, for the model to write from."""
    kind = (fmt or "").lower()
    try:
        if kind in ("excel", "xlsx", "xlsm"):
            return _xlsx(data)
        if kind == "csv":
            return _csv(data)
        if kind == "pdf":
            return _pdf(data)
        if kind in ("word", "docx"):
            return _docx(data)
        return _prose("Text file", data.decode("utf-8", errors="replace"))
    except Exception as e:  # noqa: BLE001 - an unreadable file still gets a card
        return f"(The file could not be read back: {type(e).__name__})"


def _number(value: str) -> float | None:
    v = value.replace(",", "").strip()
    return float(v) if _NUMBER.fullmatch(v) else None


def _fmt(n: float) -> str:
    return str(int(n)) if n == int(n) else f"{n:g}"


def _cells(row) -> list[str]:
    return [("" if c is None else str(c)).strip() for c in row]


_NOTES_NAME = re.compile(r"\b(summary|notes?|about|info|readme|metadata)\b", re.IGNORECASE)


def _table(name: str, raw_rows) -> str:
    rows = [r for r in (_cells(r) for r in raw_rows) if any(r)]
    if not rows:
        return f"## {name}\n(empty)"
    width = max(len(r) for r in rows)
    rows = [(r + [""] * width)[:width] for r in rows]

    # A notes sheet (a title line, a description, key: value pairs) is read as
    # notes, not data: by its name, or because its first row is no header — a
    # title with nothing beside it. A small two-column table ("Title, Year")
    # has a full header and stays data.
    header_incomplete = not all(rows[0])
    if width <= 2 and len(rows) <= 12 and (_NOTES_NAME.search(name) or header_incomplete):
        return f"## {name} (notes)\n" + "\n".join(
            f"- {r[0]}: {r[1]}" if len(r) > 1 and r[1] else f"- {r[0]}" for r in rows
        )

    header, body = rows[0], rows[1:]
    lines = [
        f"## {name}: {len(body)} data rows, {len(header)} columns",
        "It has a header row naming the columns: "
        + ", ".join(h or f"Column {i + 1}" for i, h in enumerate(header)),
        "Per column:",
    ]
    for i, col in enumerate(header):
        values = [r[i] for r in body if r[i]]
        if not values:
            continue
        numbers = [_number(v) for v in values]
        if all(n is not None for n in numbers):
            lo, hi = min(numbers), max(numbers)  # type: ignore[type-var]
            lines.append(f"- {col or f'Column {i + 1}'}: numbers from {_fmt(lo)} to {_fmt(hi)}")
            continue
        distinct = sorted(set(values))
        if len(distinct) <= _MAX_DISTINCT:
            shown = ", ".join(v[:40] for v in distinct)
            lines.append(f"- {col or f'Column {i + 1}'}: {len(distinct)} distinct values ({shown})")
        else:
            lines.append(f"- {col or f'Column {i + 1}'}: {len(distinct)} distinct values")
    # Rows are shown as "column = value" so nothing here reads like a file format.
    def show(r: list[str]) -> str:
        return "  " + "; ".join(f"{h or f'Column {i + 1}'} = {c[:40]}" for i, (h, c) in enumerate(zip(header, r)))

    lines.append("First data rows (below the header):")
    lines += [show(r) for r in body[:_SAMPLE_ROWS]]
    if len(body) > _SAMPLE_ROWS:
        lines.append("Last data row:")
        lines.append(show(body[-1]))
    return "\n".join(lines)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = wb.worksheets
        parts = [f"Excel workbook with {len(sheets)} sheet(s): " + ", ".join(ws.title for ws in sheets)]
        for ws in sheets[:_MAX_SHEETS]:
            parts.append(_table(f"Sheet '{ws.title}'", ws.iter_rows(values_only=True)))
        return "\n\n".join(parts)
    finally:
        wb.close()


def _csv(data: bytes) -> str:
    text = data.decode("utf-8-sig", errors="replace")
    rows = [r for r in csv.reader(io.StringIO(text)) if r and not r[0].startswith("# ")]
    return "CSV file\n\n" + _table("Table", rows)


def _prose(label: str, text: str) -> str:
    flat = re.sub(r"[ \t]+", " ", text)
    flat = re.sub(r"\n{3,}", "\n\n", flat).strip()
    more = " (continues)" if len(flat) > _TEXT_CHARS else ""
    return f"{label}, {len(flat.split())} words\n\n{flat[:_TEXT_CHARS]}{more}"


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    return _prose(f"PDF, {len(reader.pages)} page(s)", text)


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines = []
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        style = (p.style.name if p.style is not None else "") or ""
        lines.append(("# " if style.lower().startswith("heading") else "") + p.text.strip())
    label = f"Word document, {len(doc.tables)} table(s)"
    return _prose(label, "\n".join(lines))
