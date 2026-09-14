"""Show what the code produced, before anyone commits to it.

Generated code can succeed and still be wrong: it exits zero, writes a file,
and has quietly dropped the last twenty columns or turned every blank cell into
the text "nan". Exit status cannot catch that; opening the result can.

So the output is opened and measured, and the measurements a person would
actually check are reported next to the input's: row count, column count,
whether the columns still match, whether any cell now says "nan". The preview
rows are there so the shape is visible at a glance rather than described.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

from app.document_platform.execution.introspect import DocumentStructure

PREVIEW_ROWS = 20
MAX_CELL_CHARS = 40

_JUNK = {"nan", "none", "null", "nat", "<na>", "#n/a"}
"""Text that means a blank cell was destroyed on the way out."""


@dataclass
class Check:
    name: str
    passed: bool
    detail: str = ""


@dataclass
class ResultPreview:
    kind: str                                   # xlsx | csv | docx | pdf | other
    summary: str
    headers: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    checks: list[Check] = field(default_factory=list)
    row_count: int = 0
    column_count: int = 0

    @property
    def all_passed(self) -> bool:
        return all(c.passed for c in self.checks)

    def as_dict(self) -> dict:
        return {
            "kind": self.kind,
            "summary": self.summary,
            "headers": self.headers,
            "rows": self.rows,
            "row_count": self.row_count,
            "column_count": self.column_count,
            "checks": [
                {"name": c.name, "passed": c.passed, "detail": c.detail} for c in self.checks
            ],
            "all_passed": self.all_passed,
        }


def _cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").strip()[:MAX_CELL_CHARS]


def _spreadsheet_checks(
    headers: list[str], rows: list[list[str]], row_count: int,
    source: DocumentStructure | None, junk_cells: int,
) -> list[Check]:
    checks = [
        Check("Output opens as a valid file", True, f"{row_count} rows read back"),
        Check(
            "No NaN/None/null text in cells",
            junk_cells == 0,
            "clean" if junk_cells == 0 else f"{junk_cells} cells contain placeholder text",
        ),
        Check("Output is not empty", row_count > 0, f"{row_count} data rows"),
    ]
    if source and source.sheets:
        src = source.sheets[0]
        kept = [h for h in src.headers if h in headers]
        missing = [h for h in src.headers if h not in headers]
        checks.append(Check(
            "Original columns still present",
            not missing,
            f"{len(kept)}/{len(src.headers)} kept"
            + (f"; missing: {', '.join(missing[:5])}" if missing else ""),
        ))
        checks.append(Check(
            "Row count",
            True,
            f"{src.rows} in -> {row_count} out"
            + (" (unchanged)" if src.rows == row_count else ""),
        ))
    return checks


def _preview_excel(data: bytes, source: DocumentStructure | None) -> ResultPreview:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [_cell(h) for h in next(rows_iter)]
    except StopIteration:
        wb.close()
        return ResultPreview(kind="xlsx", summary="empty file", checks=[
            Check("Output is not empty", False, "no rows")])

    preview_rows: list[list[str]] = []
    row_count = 0
    junk = 0
    for row in rows_iter:
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        row_count += 1
        if len(preview_rows) < PREVIEW_ROWS:
            preview_rows.append([_cell(v) for v in row[: len(headers)]])
        for v in row:
            if isinstance(v, str) and v.strip().lower() in _JUNK:
                junk += 1
    sheet_names = wb.sheetnames
    wb.close()

    summary = f"{row_count} rows x {len(headers)} columns"
    if len(sheet_names) > 1:
        summary += f", {len(sheet_names)} sheets ({', '.join(sheet_names[:5])})"
    return ResultPreview(
        kind="xlsx", summary=summary, headers=headers, rows=preview_rows,
        row_count=row_count, column_count=len(headers),
        checks=_spreadsheet_checks(headers, preview_rows, row_count, source, junk),
    )


def _preview_csv(data: bytes, source: DocumentStructure | None) -> ResultPreview:
    import csv

    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    try:
        headers = [_cell(h) for h in next(reader)]
    except StopIteration:
        return ResultPreview(kind="csv", summary="empty file", checks=[
            Check("Output is not empty", False, "no rows")])
    preview_rows, row_count, junk = [], 0, 0
    for row in reader:
        if not any(str(v).strip() for v in row):
            continue
        row_count += 1
        if len(preview_rows) < PREVIEW_ROWS:
            preview_rows.append([_cell(v) for v in row[: len(headers)]])
        junk += sum(1 for v in row if v.strip().lower() in _JUNK)
    return ResultPreview(
        kind="csv", summary=f"{row_count} rows x {len(headers)} columns",
        headers=headers, rows=preview_rows, row_count=row_count,
        column_count=len(headers),
        checks=_spreadsheet_checks(headers, preview_rows, row_count, source, junk),
    )


def _preview_docx(data: bytes, _source: DocumentStructure | None) -> ResultPreview:
    import docx

    doc = docx.Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return ResultPreview(
        kind="docx",
        summary=f"{len(paragraphs)} paragraphs, {len(doc.tables)} tables",
        headers=["Text"],
        rows=[[p[:MAX_CELL_CHARS]] for p in paragraphs[:PREVIEW_ROWS]],
        row_count=len(paragraphs),
        checks=[
            Check("Output opens as a valid .docx", True),
            Check("Document is not empty", bool(paragraphs or doc.tables)),
        ],
    )


def _preview_pdf(data: bytes, _source: DocumentStructure | None) -> ResultPreview:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    text = ""
    if reader.pages:
        try:
            text = reader.pages[0].extract_text() or ""
        except Exception:  # noqa: BLE001 - a page that will not extract is still a page
            text = ""
    return ResultPreview(
        kind="pdf", summary=f"{len(reader.pages)} pages", headers=["Page 1"],
        rows=[[ln[:MAX_CELL_CHARS]] for ln in text.splitlines()[:PREVIEW_ROWS] if ln.strip()],
        row_count=len(reader.pages),
        checks=[
            Check("Output opens as a valid PDF", True),
            Check("Document has pages", len(reader.pages) > 0, f"{len(reader.pages)} pages"),
        ],
    )


def preview_output(
    data: bytes, output_name: str, source: DocumentStructure | None = None,
) -> ResultPreview:
    """Open the produced file and report what is actually in it.

    Never raises. A file that cannot be opened is the single most important
    thing to tell the user, so it is reported as a failed check rather than as
    an exception that loses the rest of the result.
    """
    ext = output_name.rsplit(".", 1)[-1].lower() if "." in output_name else ""
    readers = {
        "xlsx": _preview_excel, "xlsm": _preview_excel,
        "csv": _preview_csv, "tsv": _preview_csv,
        "docx": _preview_docx,
        "pdf": _preview_pdf,
    }
    reader = readers.get(ext)
    try:
        if reader is not None:
            return reader(data, source)
        text = data.decode("utf-8", errors="replace")
        lines = [ln for ln in text.splitlines() if ln.strip()]
        return ResultPreview(
            kind=ext or "other",
            summary=f"{len(data)} bytes, {len(lines)} lines",
            headers=["Content"],
            rows=[[ln[:MAX_CELL_CHARS]] for ln in lines[:PREVIEW_ROWS]],
            row_count=len(lines),
            checks=[Check("Output is not empty", bool(data), f"{len(data)} bytes")],
        )
    except Exception as e:  # noqa: BLE001 - an unopenable result is the headline
        return ResultPreview(
            kind=ext or "other",
            summary=f"{len(data)} bytes - could not be opened",
            checks=[Check(
                f"Output opens as a valid .{ext}", False,
                f"{type(e).__name__}: {e}",
            )],
        )
